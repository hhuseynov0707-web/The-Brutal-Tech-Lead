"""CV upload handling: text extraction and a short-lived in-memory profile store.

Only the structured profile is kept (never the raw CV), in memory, for a
limited time — nothing is written to disk.
"""

import io
import re
import secrets
import time
from collections import OrderedDict
from pathlib import Path

from .schemas import CandidateProfile

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 15_000
MIN_TEXT_CHARS = 200
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class CVError(ValueError):
    """Raised when an uploaded file cannot be used as a CV."""


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise CVError("Şifrəli PDF oxuna bilmir.")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Return normalised plain text from a PDF, DOCX, TXT or MD file. Blocking."""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise CVError("Yalnız PDF, DOCX, TXT və ya MD faylı yükləyin.")

    try:
        if suffix == ".pdf":
            text = _pdf_text(data)
        elif suffix == ".docx":
            text = _docx_text(data)
        else:
            text = data.decode("utf-8", errors="replace")
    except CVError:
        raise
    except Exception as exc:
        raise CVError("Fayl oxuna bilmədi — zədələnmiş ola bilər.") from exc

    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if len(text) < MIN_TEXT_CHARS:
        raise CVError("CV-dən kifayət qədər mətn çıxarıla bilmədi (skan edilmiş PDF ola bilər).")
    return text[:MAX_TEXT_CHARS]


class ProfileStore:
    """Tiny TTL + size-bounded store keyed by an unguessable id."""

    def __init__(self, ttl_seconds: int = 2 * 60 * 60, max_items: int = 200):
        self._ttl = ttl_seconds
        self._max = max_items
        self._items: OrderedDict[str, tuple[float, CandidateProfile]] = OrderedDict()

    def put(self, profile: CandidateProfile) -> str:
        self._evict()
        cv_id = secrets.token_urlsafe(16)
        self._items[cv_id] = (time.monotonic() + self._ttl, profile)
        while len(self._items) > self._max:
            self._items.popitem(last=False)
        return cv_id

    def get(self, cv_id: str) -> CandidateProfile | None:
        self._evict()
        item = self._items.get(cv_id)
        return item[1] if item else None

    def _evict(self) -> None:
        now = time.monotonic()
        for key in [k for k, (expires, _) in self._items.items() if expires < now]:
            del self._items[key]
